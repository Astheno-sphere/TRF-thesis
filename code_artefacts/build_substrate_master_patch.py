from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
wb=Workbook(); wb.remove(wb.active)
HF=PatternFill("solid",fgColor="0D3B66"); HFont=Font(name="Arial",bold=True,color="FFFFFF",size=10)
CF=Font(name="Arial",size=10); WR=Alignment(wrap_text=True,vertical="top")
def tab(n,h,rows,w=None):
    ws=wb.create_sheet(n[:31]); ws.append(h)
    for c in range(1,len(h)+1):
        x=ws.cell(1,c); x.fill=HF; x.font=HFont; x.alignment=WR
    for r in rows: ws.append(list(r)+[""]*(len(h)-len(r)))
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.font=CF; cell.alignment=WR
    for i,width in enumerate((w or [15,44]+[24]*(len(h)-2))[:len(h)],1):
        ws.column_dimensions[get_column_letter(i)].width=width
    ws.freeze_panes="A2"

tab("00_READ_ME",["Item","Detail"],[
["Purpose","CONSOLIDATED MASTER PATCH — supersedes Substrate_PATCH / PATCH2 / PATCH3 (Sept 2026). Copy each tab into the master sheet as one block."],
["Date","16 September 2026"],
["Chapters covered","Ch1 v2.4, Ch2 v2.4, Ch3 v1.2, Ch4 v1.3, Ch5 v1.2 — all anti-drift swept"],
["Word count (measured)","Ch1 4,712 | Ch2 11,090 | Ch3 4,446 | Ch4 6,382 | Ch5 3,154 | TOTAL ~29,780 (~66 pages at 450 w/p)"],
["Outstanding","F-02, F-03, F-05, F-06, F-10 open; F-11 substantially corroborated; Patch1 Drive upload failed (quota) — retry"],
])

tab("01_FETCH_REGISTER",["Fetch ID","Item","Status","Resolution / next step"],[
["F-01","SPUHiN D5.2 full PDF","PARTIAL","Checklists verified via Helsedirektoratet hearing attachment; full deliverable PDF still owed"],
["F-02","EDPB CEF Art-17 full report PDF","OPEN","Accessed via Matheson commentary + EDPB page. Recurring-issue COUNT deliberately removed from Ch2 rather than guessed"],
["F-03","EDPB Guidelines 02/2025 adopted status","OPEN","External reviewer claims final v2.0 adopted 7 July 2026 — NOT independently verified. Ch4 §4.9 carries the flag inline. Verify before Ch6 legal-corridor claims"],
["F-04","Aikebaier SoK full details","RESOLVED 15 Sept","Yitebeier Aikebaier, Nexum Ledger Ltd, ePrint 2026/1109. Destruction-IND equivalence PROPOSITION (not theorem). 5 open problems incl. #5 'erasure under multi-regime retention conflict' = our gap, #4 'composition with ZKP selective disclosure' = future work"],
["F-05","Haque et al. SLR final venue","OPEN","arXiv:2210.04541 cited; journal venue unconfirmed. Phase 11"],
["F-06","TEHDAS2 SPE specs deliverable number","PARTIAL","Publication during 2026 confirmed; deliverable number pending. Ch2 reference carries the flag"],
["F-07","Albalwy full PDF (113ms figure)","OPEN (optional)","Only needed if that figure is used; currently only 390ms / 216,631 gas cited"],
["F-08","EHDS implementing acts status","RESOLVED 16 Sept","Art 73(5): Commission implementing acts for SPE technical/organisational/security/interoperability requirements due BY 26 MARCH 2027. Not yet adopted. Verbatim from EUR-Lex"],
["F-09","Walonoski Synthea JAMIA citation","RESOLVED 15 Sept","JAMIA 25(3):230-238, doi 10.1093/jamia/ocx079, 10 authors — exact as cited"],
["F-10","eHOMB '12 March' — which year?","OPEN","Helsedirektoratet portfolio page states '12. mars' with no year. Ch1/Ch2 now state the date without asserting the year"],
["F-11","EUR-Lex direct text Art 73(1)(e)","SUBSTANTIALLY CORROBORATED","Full Art 73(1)(a)-(f) retrieved 16 Sept; (e) verbatim: 'the keeping of identifiable logs of access to and activities in the secure processing environment for the period necessary to verify and audit all processing operations in that environment; logs of access shall be kept for at least one year'. EUR-Lex HTML fetch returns page chrome only; direct PDF confirmation remains a Phase-11 nicety, not a blocker"],
["F-12","EHDS in-force vs applies-from","RESOLVED 16 Sept","Adopted 11 Feb 2025; OJ 5 Mar 2025; IN FORCE 26 Mar 2025; APPLIES generally from 26 Mar 2027 (Art 105); secondary-use Chapter IV incl. Arts 68 & 73 applies from 26 MAR 2029; further milestones 2031, 2035"],
])

tab("02_CORRECTIONS",["CORR","Issue","Resolution","Where"],[
["CORR-015","NIST SP 800-88 Rev 2 dated 2024","Final 26 Sept 2025; Chandramouli & Hibbard; doi 10.6028/NIST.SP.800-88r2","Ch1, Ch2, Ch4"],
["CORR-016","OUR overstatement: 'records being generated in Norwegian production piloting'","Helsedirektoratet reports no relevant patients yet; softened to 'beginning to be generated, at presently small volume'","Ch1 §1.1, Ch2 §2.2/§2.6"],
["CORR-017","eHOMB approval asserted as 'March 2026'","Year not in source; year no longer asserted","Ch1, Ch2"],
["CORR-018","Aikebaier 'equivalence theorem'","Paper says PROPOSITION; also single-author, industry-affiliated, NOT peer-reviewed — now stated","Ch2 §2.4, Ch4 §4.9"],
["CORR-019","EDPB CEF 'one of five recurring issues'","Count unverified; numeral removed rather than swapped for reviewer's 'seven'","Ch2 §2.3"],
["CORR-020","Model silent on metadata after floor expiry","Terminal transition added: destroy identifying fields + salt at t_a+F, retain non-identifiable statistical residue, log it. Art 5(1)(e)","Ch4 §4.8, §4.12"],
["CORR-021","HL7 Norway cited two ways","Harmonised to hl7.fhir.no.basis v2.2.2 (2025), Simplifier","Ch1, Ch3"],
["CORR-022","Ch3 claimed 'FHIR R4 per Norwegian basisprofiler'","Synthea default = US Core profiles, Massachusetts demographics. R4 VERSION correct, PROFILES not Norwegian. No Norwegian representativeness claimed for records","Ch3 §3.5, Ch5 §5.7"],
["CORR-023","Ch3 promised realism layer Ch5 had not built","RESOLVED BY BUILDING IT: 98 Synthea bundles, layer run, Table 5.4. All figures identical, 200/200 commitments differ","Ch3 §3.6, Ch5 §5.2/§5.7"],
["CORR-024","Patent prior art never searched","US 12,425,209 found and disclosed as Rival 5","Ch2 §2.4, Ch4 §4.10"],
["CORR-025","Invalidation entry conflated attestation with proof","Now stated: ATTESTED record of destruction, not PROOF. PoSE + certified-deletion literature marks the boundary","Ch4 §4.7, Ch2 §2.4"],
["CORR-026","EHDS timing claims imprecise","F-12 resolved: in force 2025, applies 2027, secondary use 2029. Applied throughout","Ch1 §1.1/§1.2, Ch2 §2.2/§2.3/§2.9, Ch4 §4.11"],
["CORR-027","Ch5 line count stated as 268","Measured: 322 total / 246 non-blank non-comment","Ch5 §5.2"],
])

tab("03_NEW_SOURCES",["ID","Title","Detail","Anchor","Flag"],[
["PAPER-043","Feng, Marsso, Sabetzadeh, Chechik — Bounded satisfiability checking of MFOTL","arXiv:2209.04052","Methodological anchor, Ch2 §2.5, Ch3 §3.3","G"],
["PAPER-044","Reardon, Basin, Capkun — SoK: Secure Data Deletion","IEEE S&P 2013, pp.301-315, doi 10.1109/SP.2013.28","Ch2 §2.4 honest boundary","G"],
["PAPER-045","Finck — Blockchains and data protection in the EU","EDPL 4(1):17-35 (2018)","Ch2 §2.4 residue threat","G"],
["PAPER-046","Aikebaier — SoK: Cryptographic Erasure on Public Ledgers","ePrint 2026/1109, Nexum Ledger Ltd. Preprint, not peer-reviewed","Ch2 §2.4/§2.5/§2.8/§2.9; corroborates GAP-017","G"],
["PAPER-047","Haque et al. — SLR GDPR vs public blockchain","arXiv:2210.04541","Ch2 §2.4 distinction","O"],
["PAPER-048","Chen, Chun, Patel, Chiang, James — Validity of synthetic clinical data (Synthea)","BMC Med Inform Decis Mak 19:44 (2019), doi 10.1186/s12911-019-0793-0","Ch3 §3.5 bounding","G"],
["PAPER-049","TEHDAS2 SPE technical specifications (2026)","Min capabilities + federation interoperability baseline","Ch2 §2.2/§2.6, Ch4 §4.11","O"],
["PAPER-050","US Patent 12,425,209 — Systems and methods for data retention and purging","Key-expiry selection resolves retention-vs-purge; defers purge to later date","CLOSEST PRIOR ART. Ch2 §2.4, Ch4 §4.10 Rival 5","G"],
["PAPER-051","Bursuc, Gil-Pons, Mauw, Trujillo-Rasua — Software-based memory erasure, relaxed isolation","arXiv:2401.06626 (2024)","Ch2 §2.4 outer boundary","G"],
["PAPER-052","Broadbent & Islam — Quantum encryption with certified deletion","TCC 2020, arXiv:1910.03551. Public verifiability = open problem","Ch2 §2.4; Ch7 future work","G"],
["PAPER-053","JMIR — Blockchain-based dynamic and revocable consent SLR","doi 10.2196/88536 (2026)","Ch2 adjacent consent literature","O"],
["PAPER-054","Dynamic consent for secondary use: challenges","PMC13272879 (2026). SPE logging/revocation 'not yet uniformly available'","Independent gap corroboration","O"],
["REG-016","EHDS Art 73(1)(e)","'logs of access shall be kept for at least one year' — full Art 73(1)(a)-(f) retrieved","FOURTH CONSTRAINT, absorbed as 2nd floor F=12","G"],
["REG-017","GDPR Art 12(3)","1 month, +2 extension, notice in month 1","Sets delta in C2","G"],
["REG-018","EHDS Art 73(5) + Art 105","SPE implementing acts due 26 Mar 2027; Chapter IV applies 26 Mar 2029","Ch1 §1.2, Ch2 §2.2, Ch4 §4.11, Ch6 policy timing","G"],
["NORWAY-LAW-017","Datatilsynet NIF decision","11 May 2021, NOK 1,250,000; no lawful basis to test on 3.2M records","Ch3 §3.8 ethics PRIMARY","G"],
["NORWAY-LAW-018/019","MyHealth@EU Helsedirektoratet + NHN portfolio pages","Pilot live Bodo+Stjordal; production w/ Czechia+Portugal; Trondheim+Alesund requested; NO relevant patients yet; eP w/o eD approved eHOMB 12 March","Exact URLs pinned","G"],
])

tab("04_ARTEFACTS",["Artefact","Detail","Status"],[
["trf_checker.py","322 lines / 246 code. Two record classes, both semantics, independent brute-force cross-check over 2^14 assignments. Stdlib only. Runtime 0.019s","COMPLETE — externally re-executed, byte-identical"],
["synthea_layer.py","132 lines. Substitutes real FHIR R4 payloads, recomputes commitments, compares all figures","COMPLETE — all figures invariant, 200/200 commitments differ"],
["Synthea cohort","java -jar synthea.jar -s 20260915 -cs 20260915 -p 100 --exporter.fhir.export=true ... -> 98 bundles, 314MB","COMPLETE — regenerable, not bundled"],
["synthea_sample/","3 bundles trimmed to 25 entries for inspection","COMPLETE"],
["README_Appendix_A.md","All commands, provenance, integrity table (11 results + reference hash), 5 stated limitations","COMPLETE"],
["Key results","SemI: 95 violations / 90 records (41 C2-only, 44 C3-only, 5 both). Cor 1.2: 39 witnesses, no erasure request. SemII: 0 / 0. Invalidations 138 (52 Art17, 86 Art68(12))","VERIFIED"],
])

tab("05_CHAPTER_STATE",["Chapter","Version","Words","Status"],[
["Ch1 Introduction","v2.4","4,712","COMPLETE"],
["Ch2 Literature Review","v2.4","11,090","COMPLETE — ~1,900w over budget (DEFER-05 polish debt)"],
["Ch3 Methodology","v1.2","4,446","COMPLETE"],
["Ch4 TRF Model","v1.3","6,382","COMPLETE"],
["Ch5 Demonstration","v1.2","3,154","COMPLETE"],
["Ch6 Discussion","NOT STARTED","—","6 engagements: SecuPi cost; key-destruction assurance (Reardon/PoSE); Synthea decision; timing vs 2027 acts / 2029 application; boundaries-as-asset; SEMANTICS II AUTHORITY RISK"],
["Ch7 Conclusion","NOT STARTED","—","Future work: ZKP selective disclosure (Aikebaier open problem 4); public verifiability of deletion certificates; Datatilsynet opinion on Semantics II"],
["Front/back matter","NOT STARTED","—","Abstract, declarations, consolidated references, Appendix A"],
])
wb.save("/mnt/user-data/outputs/Substrate_MASTER_PATCH_16Sept2026.xlsx"); print("master patch saved:",len(wb.worksheets),"tabs")
