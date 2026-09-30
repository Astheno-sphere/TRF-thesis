from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
wb=Workbook(); wb.remove(wb.active)
HF=PatternFill("solid",fgColor="7B2D26"); HFont=Font(name="Arial",bold=True,color="FFFFFF",size=10)
CF=Font(name="Arial",size=10); WR=Alignment(wrap_text=True,vertical="top")
def tab(n,h,rows,w=None):
    ws=wb.create_sheet(n[:31]); ws.append(h)
    for c in range(1,len(h)+1):
        x=ws.cell(1,c); x.fill=HF; x.font=HFont; x.alignment=WR
    for r in rows: ws.append(list(r)+[""]*(len(h)-len(r)))
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.font=CF; cell.alignment=WR
    for i,width in enumerate((w or [16,46]+[26]*(len(h)-2))[:len(h)],1):
        ws.column_dimensions[get_column_letter(i)].width=width
    ws.freeze_panes="A2"

tab("A_NEW_ROWS",["ID","Target Subsheet","Title / Content","Key detail","Thesis anchor","Flag","Date","Verification trail"],[
["REG-016","4 Regulations","EHDS Art 73(1)(e) — SPE identifiable access/activity logs","'logs of access shall be kept for at least one year'; collides with Art 68(12) 6-month ceiling on same environment","FOURTH CONSTRAINT — absorbed as 2nd floor instance F=12. Ch1 §1.1/§1.2/§1.3, Ch2 §2.3, Ch4 Cor 1.2","O","2026-09-15","Two independent reproductions (streamlex; metadata.healthdataportal.eu). Direct EUR-Lex text fetch still owed"],
["REG-017","4 Regulations","GDPR Art 12(3) — DSR response period","1 month from receipt, extendable +2 months if complex/numerous, notice within month 1","Sets δ in constraint C2 (Ch4 §4.4). Model uses δ=1 baseline, δ=3 extended","G","2026-09-15","Multiple independent sources incl. IAPP"],
["PAPER-048","1 Sources","The validity of synthetic clinical data: validation of Synthea using clinical quality measures","Chen J, Chun D, Patel M, Chiang E, James J. BMC Med Inform Decis Mak 19:44 (2019). doi 10.1186/s12911-019-0793-0. 1.2M patients; strong on demographics/service probability, weak on post-intervention outcomes + regional variation","Ch3 §3.5 — honest bounding of Synthea use","G","2026-09-15","Verified"],
["PAPER-049","1 Sources","TEHDAS2 SPE technical specifications (2026)","Minimum technical/functional/security capabilities for SPEs + interoperability baseline for federation; input to standardisation and implementing acts","Ch2 §2.2/§2.6, Ch4 §4.11. RESOLVES F-06 (partially — deliverable number still owed)","O","2026-09-15","Search-verified; deliverable no. pending"],
["NORWAY-LAW-017","12 Norway","Datatilsynet vedtak — Norges idrettsforbund (NIF)","11 May 2021; NOK 1,250,000 (reduced from 2.5M varsel); no lawful basis to test with 3.2M individuals' data; processing must be necessary; purpose achievable by less invasive means","Ch3 §3.8 ethics — PRIMARY source replacing Berngaard T2 commentary","G","2026-09-15","datatilsynet.no decision page"],
["NORWAY-LAW-018","12 Norway","MyHealth@EU — Helsedirektoratet portfolio page","Pilot live Bodø+Stjørdal legevakt, connected Portugal/Tsjekkia/Finland; eP uten eD change proposal approved eHOMB 12 March; legal basis+PIN+DPIA delivered 31.03.2025; NO relevant patients yet","Ch1 §1.1, Ch2 §2.2/§2.6, Ch4 §4.11 — EXACT URL now pinned","G","2026-09-15","helsedirektoratet.no/.../myhealth-eu-helsedirektoratet"],
["NORWAY-LAW-019","12 Norway","MyHealth@EU — Norsk helsenett portfolio page","Bodø+Stjørdal 'i produksjon' with Tsjekkia+Portugal; wave 8; Trondheim+Ålesund (Helseplattformen) have ordered service; NHN own solution used as vendor not ready","Ch2 §2.2 expansion detail","G","2026-09-15","helsedirektoratet.no/.../myhealth-eu-norsk-helsenett"],
])

tab("B_CORRECTIONS",["CORR ID","What was wrong","Correct position","Where fixed","Source","Flag","Date"],[
["CORR-015","NIST SP 800-88 Rev 2 cited as 2024","Final published 26 Sept 2025; authors Chandramouli & Hibbard; DOI 10.6028/NIST.SP.800-88r2; Rev 1 withdrawn same day","Ch1, Ch2, Ch4 — all instances + reference entries","csrc.nist.gov/pubs/sp/800/88/r2/final","G","2026-09-15"],
["CORR-016","OUR OWN OVERSTATEMENT: 'audit records being generated in Norwegian production piloting'","Helsedirektoratet reports NO relevant patients yet presented; infrastructure live, record volume minimal. Claim softened to 'beginning to be generated, at presently small volume'","Ch1 §1.1, Ch2 §2.2/§2.6","Helsedir portfolio page","G","2026-09-15"],
["CORR-017","eHOMB approval stated as 'March 2026'","Source says '12. mars' without year. Year NOT confirmed — flagged inline as F-10 pending","Ch1 §1.1, Ch2 §2.2","Helsedir portfolio page","O","2026-09-15"],
["CORR-018","Aikebaier SoK described as proving an 'equivalence theorem'","Paper states equivalence PROPOSITION. Also: single-author, industry-affiliated (Nexum Ledger Ltd), NOT peer-reviewed — now stated in Ch2","Ch2 §2.4, Ch4 §4.9","eprint.iacr.org/2026/1109","G","2026-09-15"],
["CORR-019","EDPB CEF 'one of five recurring issues'","Count unverified by us (external reviewer says seven). Numeral REMOVED rather than swapped — now 'among the recurring issues'","Ch2 §2.3","Precautionary","G","2026-09-15"],
["CORR-020","MODEL INCOMPLETENESS: no constraint on metadata after floor expiry","m(a) contains patient pseudonym = personal data; Art 5(1)(e) storage limitation applies at t(a)+F(a). Terminal transition added: destroy identifying fields + salt, retain non-identifiable statistical residue","Ch4 §4.8 new subsection + §4.12","Art 5(1)(e) GDPR; external review catch","G","2026-09-15"],
["CORR-021","HL7 Norway cited as GitHub repo (2024) in Ch1 vs Simplifier package (2025) in Ch3","Harmonised to hl7.fhir.no.basis v2.2.2 (2025), Simplifier — GitHub repo publishes no releases","Ch1, Ch3","simplifier.net","G","2026-09-15"],
])

tab("C_FETCH_REGISTER",["Fetch ID","Item","Status","Why it matters","Owed by"],[
["F-01","SPUHiN D5.2 full PDF","Partially resolved — checklists verified","Ch4/Ch6 citations","Ch6"],
["F-02","EDPB CEF Art-17 full report PDF","OPEN","Confirm recurring-issue count + retention findings","Ch6"],
["F-03","EDPB Guidelines 02/2025 adopted status","OPEN — external reviewer claims final v2.0 adopted 7 July 2026; NOT verified by us","Load-bearing for Ch4 §4.9 legal corridor","Ch6 / Phase 11"],
["F-04","Aikebaier SoK details","RESOLVED","Author, title, Destruction-IND, 5 open problems incl. multi-regime retention conflict","—"],
["F-05","Haque et al. SLR final venue","OPEN","Reference completeness","Phase 11"],
["F-06","TEHDAS2 SPE specs deliverable number","PARTIAL — publication confirmed, number pending","Ch2/Ch4 citation precision","Ch6"],
["F-07","Albalwy full PDF","OPEN (optional)","Only if 113ms figure used","Optional"],
["F-08","EHDS implementing acts status","PARTIAL — due 26 Mar 2027, not adopted","Ch6 policy claims","Ch6"],
["F-09","Walonoski Synthea JAMIA citation","RESOLVED","25(3):230-238, doi 10.1093/jamia/ocx079, exact","—"],
["F-10","eHOMB '12 March' — which year?","NEW / OPEN","Ch1+Ch2 assert approval; year unconfirmed","Before submission"],
["F-11","EUR-Lex direct text of Art 73(1)(e)","NEW / OPEN","Two secondary reproductions only; provision is load-bearing for Cor 1.2","Before Ch5"],
])

tab("D_CHAPTER_STATE",["Chapter","Status","Words","Notes"],[
["Ch1 Introduction","v2.2 COMPLETE","~3,300","RQ generalised to audit-retention floors; NCPeH precision; NIST fix; Aikebaier corroboration"],
["Ch2 Literature Review","v2.2 COMPLETE","~8,800","Art 73(1)(e) in §2.3; TEHDAS2 2026 currency; Gap 1 anchored to SoK open-problem list"],
["Ch3 Methodology","v1.1 COMPLETE","~3,700","Synthea validation bounded; Datatilsynet primary; two record classes in checker"],
["Ch4 TRF Model","v1.1 COMPLETE","~6,400","Theorems 1+2 with proofs; Cor 1.2 = no-erasure-request infeasibility; residue analysis; end-of-life transition"],
["Ch5 Demonstration","NOT STARTED","—","BUILD CHECKER FIRST, then write around real output"],
["Ch6 Discussion","NOT STARTED","—","Requires Ch5 results + F-02/F-03/F-08 closed"],
["Ch7 Conclusion","NOT STARTED","—",""],
["Front/back matter","NOT STARTED","—","Phase 11"],
])
wb.save("/mnt/user-data/outputs/Substrate_PATCH2_Sept2026.xlsx")
print("saved", len(wb.worksheets),"tabs")
