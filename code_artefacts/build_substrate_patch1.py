# -*- coding: utf-8 -*-
"""Substrate PATCH workbook - Sept 2026. All rows owed to the master sheet."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook(); wb.remove(wb.active)
HF = PatternFill("solid", fgColor="1F3864"); HFont = Font(name="Arial", bold=True, color="FFFFFF", size=10)
CF = Font(name="Arial", size=10); WR = Alignment(wrap_text=True, vertical="top")

def tab(name, hdr, rows, w=None):
    ws = wb.create_sheet(name[:31]); ws.append(hdr)
    for c in range(1, len(hdr)+1):
        x = ws.cell(1, c); x.fill = HF; x.font = HFont; x.alignment = WR
    for r in rows:
        ws.append(list(r) + [""]*(len(hdr)-len(r)))
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.font = CF; cell.alignment = WR
    for i, width in enumerate((w or [16,42]+[24]*(len(hdr)-2))[:len(hdr)], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

S16 = ["Source ID","Title","Authors","Year","Type","Source/Journal","DOI/URL","Database","Keywords","Relevance Tag","Thesis Anchor","Verification","Flag","Flag-Date","Flag-Note","Verification-Trail"]

tab("1_Sources_ADD", S16, [
["PAPER-031","Developing an infrastructure for secure patient summary exchange in the EU context: Lessons learned from the KONFIDO project","Natsiavas P, Mazzeo G, Faiella G, Campegiani P, Dumortier J, Stan O, et al. (17 authors, Gelenbe E, Romano L senior)","2021","Peer-reviewed","Health Informatics Journal 27(2)","10.1177/14604582211021459","Sage","KONFIDO, OpenNCP, cross-border, lessons learned","KONFIDO lessons-learned — distinct from PAPER-019 (2018 RG item) and PAPER-036 (2019 Springer)","Ch1 + Ch2 §2.4 cited (CITE-005)","Verified-primary","G","2026-09-15","Fills missing row; Chapter 1/2 cite this as separate source from the two earlier KONFIDO papers","Web search verified 2026-05-22"],
["PAPER-043","Bounded Satisfiability Checking of Metric First-Order Temporal Logic","Feng N, Marsso L, Sabetzadeh M, Chechik M","2022","Peer-reviewed (ICSE-track)","arXiv","arXiv:2209.04052","arXiv","MFOTL, SMT, Z3, legal compliance, satisfiability, counterexamples","Formal-verification-of-compliance anchor","METHODOLOGICAL ANCHOR for TRF Model (Ch2 §2.5, Ch3, Ch4)","Verified-primary","G","2026-09-15","Novelty-audit decisive find: method exists, tri-lateral health-audit application does not","Web search Jul 2026 novelty audit"],
["PAPER-044","SoK: Secure Data Deletion","Reardon J, Basin D, Capkun S","2013","Peer-reviewed canonical","2013 IEEE Symposium on Security and Privacy, pp. 301-315","10.1109/SP.2013.28","IEEE","secure deletion, sanitization, adversary taxonomy, irrecoverability","Canonical secure-deletion systematisation; 'further from medium = harder to assert destruction'","Ch2 §2.4 honest boundary of crypto-erasure; Ch6 key-destruction limitation","Verified-primary","G","2026-09-15","ETH Zurich; verified via Semantic Scholar + ETH syssec BibTeX","Web search 2026-09-15"],
["PAPER-045","Blockchains and data protection in the European Union","Finck M","2018","Peer-reviewed","European Data Protection Law Review 4(1), 17-35","EDPL 4(1):17-35","EDPL","blockchain, GDPR, immutability, right to be forgotten","Canonical academic statement of immutability-vs-RTBF; hashed data likely personal data","Ch2 §2.4 adjacent-literature distinction + residue threat (drives Ch4 hiding commitments)","Verified-primary","G","2026-09-15","Verified via RWA-tokenization SoK reference list + EP study lineage","Web search 2026-09-15"],
["PAPER-046","SoK: Cryptographic Erasure on Public Ledgers","IACR ePrint authors (full list pending PDF read)","2026","Preprint (systematisation)","Cryptology ePrint Archive 2026/1109","eprint.iacr.org/2026/1109","IACR","cryptographic erasure, Destruction-IND, key lifecycle, EDPB 02/2025, Article 17","Formalises the ERASURE side (equivalence theorem to EDPB unrecoverability test); explicitly excludes retention under sectoral law","Ch2 §2.4/§2.5/§2.8 nearest analytical neighbour; GAP-017 survives by its own scope statement","Verified-primary (abstract+intro)","O","2026-09-15","Full PDF read owed at Ch4 prep: author list + Destruction-IND definition details","Direct eprint fetch snippet 2026-09-15"],
["PAPER-047","A Systematic Literature Review of the Tension between the GDPR and Public Blockchain Systems","Haque AB, et al.","2022","SLR (arXiv version)","arXiv 2210.04541","arXiv:2210.04541","arXiv","GDPR, blockchain, RTBF, immutability, SLR","RTBF-immutability = most-discussed topic (53 papers) — evidences the ADJACENT conflict is heavily studied, ours is distinct","Ch2 §2.4 distinction: technological vs regulatory conflict","Verified-primary","O","2026-09-15","Final journal venue + full author list to confirm at Phase 11","Web search 2026-09-15"],
])

tab("4_Regs_ADD", S16, [
["REG-015","EDPB Guidelines 02/2025 on the processing of personal data through blockchain technologies","European Data Protection Board","2025","Primary regulatory guidance","EDPB","edpb.europa.eu (Guidelines 02/2025)","EDPB","blockchain, erasure, cryptographic unrecoverability, threat model, Article 17","DECISIVE: accepts rendering-unrecoverable-through-cryptographic-means as erasure substitute IF threat model documented","Ch1 §1.4 + Ch2 §2.3 legal corridor; Ch4 designs to this test","Verified via IACR SoK citation + commentary","O","2026-09-15","Adopted-version status (consultation vs final) to verify at Phase 11 — cite criterion carefully until then","Web search 2026-09-15"],
])

N13 = ["Source ID","Title (NO)","Title (EN gloss)","Issuing Body","Document Type","Date","Status","URL","Relevance to Thesis","Flag","Flag-Date","Flag-Note","Verification-Trail"]
tab("12_Norway_ADD", N13, [
["NORWAY-LAW-013","MyHealth@EU-prosjektet — pilot og fase 2","Norwegian NCPeH pilot status","Helsedirektoratet","Official project page","Read Sept 2026","Live","helsedirektoratet.no (MyHealth@EU project page)","Pilot LIVE at Bodø + Stjørdal legevakt with Portugal/Czechia/Finland; 2026 phase 2 = NCP-A; eP WITHOUT eD unique among states, eHOMB approval 12 March; prod test Q2 2026. Upgrades Ch1 timeliness","G","2026-09-15","Resolves DEFER-06 NCPeH item","Web search 2026-09-15"],
["NORWAY-LAW-014","Anonymisering av personopplysninger — veileder (2015)","Datatilsynet anonymisation guide","Datatilsynet","DPA guidance (veileder)","2015","In force","datatilsynet.no/.../anonymisering-veileder-041115.pdf","HIGH anonymity threshold: 'kryptering er å beskytte data, ikke å gjøre dem uidentifiserbare'; key/linkage = still personal. Constrains Theorem 2 residue analysis","G","2026-09-15","Constraining side of the legal corridor","Direct retrieval 2026-09-15"],
["NORWAY-LAW-015","Etablere internkontroll","Datatilsynet internal-control guidance","Datatilsynet","DPA guidance","Current","Live","datatilsynet.no/rettigheter-og-plikter/virksomhetenes-plikter/informasjonssikkerhet-internkontroll/etablere-internkontroll/","Controllers must clarify 'lovlig oppbevaringstid og slettefrister' weighing 'oppbevaringsrett og oppbevaringsplikt' — Norwegian DPA names our retention-vs-deletion balancing as controller duty","G","2026-09-15","Supports Ch2 §2.3 Art 17(3)(b) defence (Datatilsynet n.d.-a)","Direct retrieval 2026-09-15"],
["NORWAY-LAW-016","Innebygd personvern og personvern som standard","Privacy-by-design guidance","Datatilsynet","DPA guidance","Current","Live","datatilsynet.no/rettigheter-og-plikter/virksomhetenes-plikter/innebygd-personvern-og-personvern-som-standard/","'Anonymisering likestilles med sletting av personopplysninger' — anonymisation EQUATED with deletion. Supporting side of legal corridor for Theorem 2","G","2026-09-15","Cited as Datatilsynet n.d.-b in Ch2 §2.3","Direct retrieval 2026-09-15"],
["SPUHIN-DOC-004","D5.2 Recommendations for minimum requirements for Secure Processing (final)","Final SPE criteria","SPUHiN / Helsedirektoratet + FHI","Government deliverable (final)","Hearing Jul-Oct 2025; published after","Published","helsedirektoratet.no SPUHiN page + helsedirektoratet.no hearing-attachment PDF (checklists)","Final checklists (P1 ISMS Mandatory, P3 ISO27001 Recommended, eID/BankID levels, personnel security); national reference basis pending EU implementing act; NO retention-erasure mechanism at any grade — GAP SURVIVES FINAL DOCUMENT","G","2026-09-15","Upgraded from O: verified via Helsedirektoratet project page + checklist attachment content","Direct retrieval 2026-09-15"],
])

P11 = ["Policy ID","Title","Issuer","Date","Tier","URL","Key Content","Thesis Anchor","Flag","Flag-Date","Note"]
tab("19_Policy_NEW", P11, [
["TIER-DEF","T1 = official policy study/intergovernmental; T2 = professional/legal commentary; T3 = industry/vendor commentary (context + counter-positions only, never load-bearing for factual regulatory claims)","—","—","—","—","Schema row — keep at top of Subsheet 19","—","G","2026-09-15","Create this tab as Subsheet 19: Secondary & Policy Sources"],
["POLICY-001","Data Flows in Supply Chains policy brief","International Chamber of Commerce","2025 (Jun; Feb 2026 update)","T1","iccwbo.org","Cross-border data flows structurally equivalent to material flows; OECD: data autarky = -4.5% GDP, -8.5% exports","Ch2 §2.7 information-logistics framing (LSL-029)","G","2026-09-15","Was drafted May 2026, never pasted"],
["POLICY-002","EHDS: Europe's Health Data Space update","Healthy Europe","2026","T2","healthyeurope.eu/european-health-data-space/","Implementation calendar tracking; TEHDAS2 29 countries to Dec 2026","Ch2 §2.2 calendar specifics","Y","2026-09-15","Secondary tracking source"],
["POLICY-003","Regulation (EU) 2025/327 EHDS alert","Ernst & Young (Greece)","2025","T2","ey.com","National Digital Health Authority by Jun 2025; provider/vendor certification by Jan 2026","Ch2 §2.2","Y","2026-09-15","Secondary commentary"],
["POLICY-004","EDPB publishes CEF report on the right to erasure","Matheson LLP","2026 (Mar)","T2","matheson.com/insights/edpb-publishes-cef-report-on-the-right-to-erasure/","Commentary route into EDPB CEF findings incl. Issue 5 retention periods","Ch2 §2.3 (LSL-028 companion)","Y","2026-09-15","Full EDPB report fetch still owed (Pending_Fetches)"],
["POLICY-005","Crypto-Shredding is NOT Nirvana for RTBF Compliance","SecuPi","2025 (Jun)","T3","secupi.com/crypto-shredding-is-not-nirvana-for-right-of-erasure-or-rtbf-compliance/","Operational-complexity critique; PBAC/ABAC alternative","Essential counter-position — Ch2 §2.3 named, Ch3 §3.7 signposted, Ch6 engaged","G","2026-09-15","Academic seriousness anchor"],
["POLICY-006","Blockchain and the GDPR: Can distributed ledgers be squared with European data protection law?","Finck M — European Parliament STOA/EPRS","2019 (Jul), 120pp","T1","europarl.europa.eu (EPRS study)","Erasure/anonymisation legally uncertain; hashed/encrypted data likely remain personal; permissioned easier; recommends EDPB guidance (which became 02/2025 lineage)","Ch2 §2.4 residue threat → Ch4 hiding commitments","G","2026-09-15","T1 despite commentary style: official EP study"],
["POLICY-007","Crypto-Shredding: Reconciling GDPR and Financial Record-Keeping (MiFID II)","VeritasChain Standards Organization","2026 (18 Jan)","T3","veritaschain.org/blog/posts/2026-01-18-crypto-shredding-gdpr-mifid-ii-reconciliation/","BI-LATERAL reconciliation asserted industry-wide (floor vs right); no model, no verifiability semantics, no ceiling","Ch2 §2.4/§2.8 practitioner-pattern positioning — our tri-lateral delta","G","2026-09-15","SAVE LOCAL COPY (blog volatility)"],
["POLICY-008","Ivaretar din teknologiske tjeneste kravet til innebygd personvern?","Berngaard (NO law firm)","2024","T2","berngaard.no","Deletion must cover all media incl. backups AND LOGS; Datatilsynet case practice (Skolemelding, NIF synthetic-data point)","Ch2 §2.3 logs-in-scope sentence","Y","2026-09-15","Also useful: NIF case = DPA endorsing synthetic data for testing (Ch3 ethics support)"],
["POLICY-009","How to set up crypto-shredding for GDPR right-to-erasure compliance","Oneuptime","2026 (Feb)","T3","oneuptime.com","DPA-recognition claim for crypto-erasure in cloud, conditional on auditable key mgmt","Ch2 §2.3 (already cited v1)","Y","2026-09-15","Keep T3 weight only"],
])

G10 = ["Row Type","ID","Target Subsheet","Content / OLD","Content / NEW or Description","Rationale / Evidence","Status","Flag","Flag-Date","Note"]
tab("Gaps_Reframes_Checkpts", G10, [
["GAP","GAP-017","7 Research Gap Log","—","No formal model of the TRI-LATERAL retention system (eHDSI 10-yr floor ∧ GDPR Art 17 right ∧ Art 68(12) 6-month ceiling) in ANY sector; OR/logistics methods unapplied to data-retention reconciliation","3 novelty searches Jul 2026 + Sept verification: Feng 2022 owns method not application; IACR 2026/1109 formalises erasure side and EXCLUDES retention law by own scope; VeritasChain asserts bi-lateral only","Live — PRIMARY THESIS GAP (supersedes GAP-001U as headline)","P","2026-09-15","TRF Model occupies this"],
["REFRAME","REF-013","15 Reframing Inventory","Apply recognised crypto-erasure + accumulators to EHDS audit conflict in Norway (REF-011 application novelty)","Formally characterise the tri-lateral conflict as feasibility problem (TRF, after Feng 2022) in information-logistics terms; crypto-erasure-by-invalidation = DERIVED constructive witness, demonstrated executably; deltas vs literature = tri-lateral model + verifiability semantics + cross-border health/Norwegian instantiation","Survived 3 adversarial searches + Sept recheck (IACR SoK, VeritasChain found and positioned)","Locked — supervisor-backed","P","2026-09-15","REF-011 stands underneath for mechanism layer"],
["CHECKPOINT","v4","11 Versioned Checkpoints","—","23 May 2026: full-schema discipline enforced after abbreviation catch; substrate audit round","—","Logged","G","2026-09-15",""],
["CHECKPOINT","v5","11 Versioned Checkpoints","—","23 May 2026: GPT-doc verification round; CORR-005/006/007 (HAP avviklet, e-helse dissolution, 2024/1689)","—","Logged","G","2026-09-15",""],
["CHECKPOINT","v6","11 Versioned Checkpoints","—","24 May 2026: triple-regulation tracks + Norwegian AI-Act actors (Nkom/Digdir/Datatilsynet/Norsk Akkreditering)","—","Logged","G","2026-09-15",""],
["CHECKPOINT","v7","11 Versioned Checkpoints","—","24 May 2026: substrate locked at 17 subsheets; Phase 0-3 locks; Ch1-Ch2 v1 written","—","Logged","G","2026-09-15",""],
["CHECKPOINT","v8","11 Versioned Checkpoints","—","15 Sept 2026: resume; TRF gate passed; novelty upgrade; Ch1 v2.1 + Ch2 v2 FINAL; this patch","—","Logged","G","2026-09-15",""],
["ARTEFACT-MAP","TRF","9 Artefact Mapping (ADD ROW AT TOP)","—","TRF — Tri-Lateral Retention Feasibility Model | PRIMARY ANALYTICAL ARTEFACT | anchors: PAPER-043, GAP-017, REF-013 | Ch4 core; checker code = executable witness","Gate passed 15 Sept 2026","Live","P","2026-09-15","DEFER-07 execution part 1"],
["ARTEFACT-MAP","A3 edit","9 Artefact Mapping (EDIT Flag-Note)","(current note)","PRIMARY THESIS MECHANISM — derived from TRF. Numbering note: decision log once said 'A4'; sheet A3 = canonical from here","—","Edit","P","2026-09-15","DEFER-07 part 2"],
["ARTEFACT-MAP","A1/A2/A4/A5/A6 edit","9 Artefact Mapping (EDIT Flag-Notes)","(current notes)","ADDENDUM-TO-JOÃO — future-work signal, out of thesis scope; no SNARKs in thesis","—","Edit","P","2026-09-15","DEFER-07 part 3"],
["STALE-EDIT","STD-015","6 Technical Standards (EDIT relevance cell)","Core cryptographic primitive for VERA-SNARK","Excluded from thesis scope (no SNARKs in mechanism); addendum only; name 'VERA-SNARK' retained with design-evolution narrative (Ch1 §1.6)","NORS 2025 name-evolution decision","Edit","P","2026-09-15",""],
["STALE-EDIT","AB-002 / AB-003","3 Annotated Bibliography (PREPEND to annotations)","(current)","SUPERSEDED SCOPE — zk-SNARK framing predates TRF reframe; retain for addendum only","—","Edit","P","2026-09-15",""],
])

tab("LSL_Renumber", ["Instruction","Old ID","New ID","Row identity (so you edit the right one)"], [
["Subsheet 2 contains DUPLICATE IDs LSL-021..026. Keep the FIRST set (SPUHiN/TEHDAS2 rows) as-is. Renumber the SECOND set (post-Ch2 batch) as follows:","","",""],
["Renumber","LSL-021 (2nd)","LSL-027","REG-010 IHE ATNA row"],
["Renumber","LSL-022 (2nd)","LSL-028","REG-011 EDPB CEF Feb 2026 row"],
["Renumber","LSL-023 (2nd)","LSL-029","REG-012 NIST SP 800-88 row"],
["Renumber","LSL-024 (2nd)","LSL-030","PAPER-037 Ren accumulator survey row"],
["Renumber","LSL-025 (2nd)","LSL-031","POLICY-001 ICC data-flows row"],
["Renumber","LSL-026 (2nd)","LSL-032","POLICY-005 SecuPi row"],
])

tab("Pending_Fetches", ["Fetch ID","Item","Why needed","Deadline (phase)","Status"], [
["F-01","SPUHiN D5.2 full PDF (EU Funding & Tenders portal version)","Checklist verified via Helsedirektoratet attachment; full deliverable for Ch4/Ch6 citations","Ch4 prep (W3)","Partially resolved — checklist content verified 2026-09-15"],
["F-02","EDPB CEF Art-17 full report PDF","Currently via Matheson commentary + EDPB page","Ch6 prep","Open"],
["F-03","EDPB Guidelines 02/2025 — adopted-version status + exact unrecoverability wording","Load-bearing for legal corridor","Phase 11 latest; ideally Ch4","Open"],
["F-04","IACR ePrint 2026/1109 full PDF — author list + Destruction-IND definition","Ch2 cites; Ch4 consumes erasure semantics","Ch4 prep (W3)","Open"],
["F-05","Haque et al. SLR final journal venue + authors","Reference completeness","Phase 11","Open"],
["F-06","TEHDAS2 M7.4 SPE technical specs","Ch4 SPE integration section","Ch4 prep (W3)","Open"],
["F-07","Albalwy 2026 full PDF (113ms proof-gen figure)","Only if figure used; else cite 390ms/216,631 only","Optional","Open"],
["F-08","EHDS implementing-acts status (search skipped 15 Sept)","Ch6 policy implications currency","Ch6 prep (W7)","Open"],
["F-09","Walonoski et al. 2018 Synthea JAMIA citation cross-check","Ch3 cites canonically; verify DOI/pages","Phase 11","Open"],
])

wb.save("/home/claude/Substrate_PATCH_Sept2026.xlsx")
import os; print("saved", os.path.getsize("/home/claude/Substrate_PATCH_Sept2026.xlsx"), "bytes,", len(wb.worksheets), "tabs")
