from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
wb=Workbook(); wb.remove(wb.active)
HF=PatternFill("solid",fgColor="1B5E20"); HFont=Font(name="Arial",bold=True,color="FFFFFF",size=10)
CF=Font(name="Arial",size=10); WR=Alignment(wrap_text=True,vertical="top")
def tab(n,h,rows,w=None):
    ws=wb.create_sheet(n[:31]); ws.append(h)
    for c in range(1,len(h)+1):
        x=ws.cell(1,c); x.fill=HF; x.font=HFont; x.alignment=WR
    for r in rows: ws.append(list(r)+[""]*(len(h)-len(r)))
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.font=CF; cell.alignment=WR
    for i,width in enumerate((w or [16,46]+[26]*(len(h)-2))[:len(h)],1):
        ws.column_dimensions[get_column_letter(i)].width=width
    ws.freeze_panes="A2"

tab("A_NEW_SOURCES",["ID","Target Subsheet","Title","Detail","Thesis anchor","Flag","Date"],[
["PAPER-050","1 Sources","US Patent 12,425,209 — Systems and methods for data retention and purging","Resolves retention-vs-purge conflict by selecting encryption key whose expiration matches governing date, per-value granularity. CLOSEST PRIOR ART to our mechanism","Ch2 §2.4 prior-art disclosure; Ch4 §4.10 Rival 5. Defeated by ceiling: deferring purge is what Art 68(12) forbids","G","2026-09-16"],
["PAPER-051","1 Sources","Software-Based Memory Erasure with relaxed isolation requirements","Bursuc S, Gil-Pons R, Mauw S, Trujillo-Rasua R (2024). arXiv:2401.06626. Formal PoSE models without isolation assumption","Ch2 §2.4 outer boundary; Ch4 attested-vs-proven distinction; Ch6 limitation","G","2026-09-16"],
["PAPER-052","1 Sources","Quantum encryption with certified deletion","Broadbent A, Islam R (2020), TCC 2020, arXiv:1910.03551. Deletion certificates; PUBLIC VERIFIABILITY named as open problem","Ch2 §2.4; Ch7 future work — the assurance gap our mechanism cannot close","G","2026-09-16"],
["PAPER-053","1 Sources","Blockchain-Based Dynamic and Revocable Consent for Secondary Health Data Use: Systematic Review","JMIR Med Inform 2026, doi 10.2196/88536. Hybrid on/off-chain; hashed refs + metadata on ledger","Ch2 §2.4 adjacent consent literature (distinct from retention conflict)","O","2026-09-16"],
["PAPER-054","1 Sources","Dynamic Consent for Secondary Use of Health Data: Challenges","PMC13272879 (2026). Notes SPE logging/auditability/revocation propagation 'not yet uniformly available'; pseudonymisation impedes targeted erasure","SUPPORTS GAP — independent confirmation SPE logging unresolved. Ch2 §2.6, Ch6","O","2026-09-16"],
])

tab("B_CORRECTIONS",["CORR ID","Issue","Resolution","Where","Flag","Date"],[
["CORR-022","Ch3 claimed records 'FHIR R4 per Norwegian basisprofiler'","Synthea default export = US Core profiles, Massachusetts demographics. R4 VERSION correct, PROFILES not Norwegian. Corrected in Ch3 §3.5 and Ch5 §5.7; no Norwegian representativeness claimed for records","Ch3 §3.5/§3.6, Ch5 §5.2/§5.7","G","2026-09-16"],
["CORR-023","Ch3/Ch5 described different artefacts (realism layer promised, not built)","RESOLVED BY BUILDING IT. Synthea 98 bundles generated, layer run, Table 5.4 added. Every figure identical, 200/200 commitments differ","Ch3 §3.6, Ch5 §5.2/§5.7","G","2026-09-16"],
["CORR-024","Prior art not searched in patent literature","US 12,425,209 found and disclosed. Added as Rival 5 in Ch4 §4.10 and prior-art passage in Ch2 §2.4","Ch2 §2.4, Ch4 §4.10","G","2026-09-16"],
["CORR-025","Invalidation entry described without distinguishing attestation from proof","Ch4 §4.7 now states the entry is an ATTESTED record of destruction, not a PROOF; PoSE/certified-deletion literature marks the boundary","Ch4 §4.7, Ch2 §2.4","G","2026-09-16"],
["CORR-026","EHDS application date ambiguity surfaced in literature check","Sources variously give 'applies from 26 March 2026' and staggered Art 105 dates vs our 'in force 26 March 2025'. In-force vs applies-from distinction must be verified before Ch6 policy claims","Ch1 §1.1, Ch2 §2.2 — NOT yet changed","O","2026-09-16"],
])

tab("C_FETCH_REGISTER",["Fetch ID","Item","Status","Owed by"],[
["F-02","EDPB CEF Art-17 full report PDF","OPEN","Ch6"],
["F-03","EDPB Guidelines 02/2025 adopted status (external reviewer claims final v2.0, 7 Jul 2026 — UNVERIFIED by us)","OPEN","Ch6"],
["F-05","Haque et al. SLR final venue","OPEN","Phase 11"],
["F-06","TEHDAS2 SPE specs deliverable number","PARTIAL","Ch6"],
["F-08","EHDS implementing acts status","PARTIAL — due 26 Mar 2027","Ch6"],
["F-10","eHOMB '12 March' — which year?","OPEN","Before submission"],
["F-11","EUR-Lex direct text Art 73(1)(e)","OPEN — two independent secondary reproductions held","Before Ch6"],
["F-12","EHDS in-force vs applies-from dates (Art 105 staggering)","NEW / OPEN — surfaced 16 Sept literature check","Before Ch6"],
["F-01","SPUHiN D5.2 full PDF","PARTIAL — checklists verified","Ch6"],
["F-04","Aikebaier SoK","RESOLVED","—"],
["F-09","Walonoski Synthea citation","RESOLVED","—"],
])

tab("D_ARTEFACT_REGISTER",["Artefact","Location","Status","Verified"],[
["trf_checker.py (322 lines)","Appendix A / code_artefacts","COMPLETE — runs 0.019s, deterministic","Executed; reproduced byte-for-byte by external reviewer"],
["synthea_layer.py (132 lines)","Appendix A / code_artefacts","COMPLETE — payload-provenance comparison","Executed 16 Sept; all figures invariant, 200/200 commitments differ"],
["Synthea cohort (98 FHIR R4 bundles)","regenerable: java -jar synthea.jar -s 20260915 -cs 20260915 -p 100","COMPLETE — 314MB, not bundled","Generated 16 Sept"],
["synthea_sample/ (3 trimmed bundles)","Appendix A / code_artefacts","COMPLETE — inspection copies","—"],
["README_Appendix_A.md","Appendix A","COMPLETE — commands, provenance, integrity table, 5 stated limitations","—"],
["Substrate patch workbooks 1-3","outputs","Patch1 Drive upload FAILED (quota) — retry owed","—"],
])

tab("E_CHAPTER_STATE",["Chapter","Status","Notes"],[
["Ch1 Introduction","v2.3","+Aikebaier corroboration, NCPeH precision, RQ=floors plural, NIST 2025"],
["Ch2 Literature Review","v2.3","+prior-art patent disclosure, +PoSE/certified-deletion boundary, +Art 73(1)(e)"],
["Ch3 Methodology","v1.2","+Synthea validation bounds, +US Core profile limitation stated"],
["Ch4 TRF Model","v1.2","+Rival 5 (patent), +attested-vs-proven, +end-of-life transition"],
["Ch5 Demonstration","v1.2","+Table 5.4 Synthea comparison, +parameter table, +decomposition"],
["Ch6 Discussion","NOT STARTED","6 engagements incl. Semantics II authority risk"],
["Ch7 Conclusion","NOT STARTED","Future work: ZKP selective disclosure (Aikebaier open problem 4), public verifiability of deletion certificates, DPA opinion on Semantics II"],
])
wb.save("/mnt/user-data/outputs/Substrate_PATCH3_Sept2026.xlsx"); print("patch3 saved",len(wb.worksheets),"tabs")
